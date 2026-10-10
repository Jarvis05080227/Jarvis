# -*- coding: utf-8 -*-
"""[COWORK-VERIFY-W1-S1 / 2026-07-09] 真验收断言层 (R10 / COWORK-TRUTH-S3).

缘起 (COWORK-VERIFY 设计窗): R10 开三闸只买"诚实" (如实标 inconclusive /
cannot_tell), 买不到"真相"。本模块买真相 — 委托完工时按产物类型**真打开文件
机检**断言是否成立 (改表案: 真开 xlsx 查金额列 == 9999)。

形态 (准则 6 三件套):
  1. 持久化: op 集 + kind 集在 memory_pool/verify_assert_vocab.json (enabled 默 0)
  2. CLI: scripts/verify_assert_dump.py (list/enable/disable/--test 干跑)
  3. 消费: jarvis_cowork._verify_executor_task 一档插入 (W2 接线; 本 S1 无消费者)

红线 (对齐设计 §3.2/§3.6):
  - verdict 只 pass(真验收) / fail(收紧不销账) / cannot_tell(退兜底);
    任何异常/断供一律 cannot_tell — 只收紧不放松, 绝不因验证器故障阻塞
    既有存在性 verified 链。
  - op / reader 全用具名小函数表 (_OPS / _VERIFIERS), 新 op/新格式 = 新小函数,
    禁做通用断言表达式解析器 (准则 10, FLOWVITAL _COMP_SAMPLERS 教训)。
  - reader read-only 打开, 零写零执行 (openpyxl read_only=True 等)。
  - [W1-S3] 每 reader 打开文件前过 jarvis_file_access_policy.check_path
    (fail-toward-DENY: policy 拒/断供 → cannot_tell 'policy_denied:...')。
"""
from __future__ import annotations

import json
import os
import re
import threading
import time  # noqa: F401
from typing import Any, Dict, List, Optional, Tuple
# [CFG-SPLIT-1 / 2026-08-16] 入库 seed + 运行态覆盖层 (memory_pool/_local/,
# 出仓)。_cfg_overlay.json_load 从文件对象的 .name 取路径再合并;
# 没有覆盖层的文件是严格 no-op, 与改动前同一条路。
import jarvis_config_overlay as _cfg_overlay

_ROOT = os.path.dirname(os.path.abspath(__file__))
_VOCAB_PATH = os.path.join(_ROOT, 'memory_pool', 'verify_assert_vocab.json')

_CACHE_LOCK = threading.RLock()
_VOCAB_CACHE: Optional[dict] = None
_VOCAB_MTIME: float = 0.0

_SEED_VOCAB: Dict[str, Any] = {
    '_meta': {
        'schema': 'verify_assert_vocab',
        'schema_version': 1,
        'marker': 'COWORK-VERIFY-W1-S1 / 2026-07-09',
        'edit_via': 'scripts/verify_assert_dump.py',
    },
    'enabled': 0,
    'size_cap_bytes': 52428800,
    'ops': ['all_equals', 'equals', 'count_gte', 'contains', 'nonempty'],
    'kinds': ['xlsx', 'csv', 'json', 'sqlite', 'text'],
}


# ---------------------------------------------------------------------------
# vocab loader (mtime cache, 照搬 jarvis_cowork.py:127-149 范式)
# ---------------------------------------------------------------------------

def _load_vocab_raw() -> dict:
    if not os.path.exists(_VOCAB_PATH):
        return dict(_SEED_VOCAB)
    try:
        with open(_VOCAB_PATH, 'r', encoding='utf-8') as f:
            data = _cfg_overlay.json_load(f)
        return data if isinstance(data, dict) else dict(_SEED_VOCAB)
    except Exception:
        return dict(_SEED_VOCAB)


def _load_vocab() -> dict:
    global _VOCAB_CACHE, _VOCAB_MTIME
    try:
        # [OVERLAY-MTIME-SWEEP / 2026-09-05] 覆盖层 mtime 进缓存键 (CLI 翻 _local/ 不重启即热载)
        mtime = max(os.path.getmtime(_VOCAB_PATH), _cfg_overlay.overlay_mtime(_VOCAB_PATH)) if os.path.exists(_VOCAB_PATH) else 0.0
    except OSError:
        mtime = 0.0
    with _CACHE_LOCK:
        if _VOCAB_CACHE is not None and mtime == _VOCAB_MTIME:
            return _VOCAB_CACHE
        _VOCAB_CACHE = _load_vocab_raw()
        _VOCAB_MTIME = mtime
        return _VOCAB_CACHE


def is_verify_assert_enabled() -> bool:
    try:
        return bool(_load_vocab().get('enabled'))
    except Exception:
        return False


def get_size_cap_bytes() -> int:
    try:
        cap = int(_load_vocab().get('size_cap_bytes', _SEED_VOCAB['size_cap_bytes']))
        return cap if cap > 0 else _SEED_VOCAB['size_cap_bytes']
    except Exception:
        return _SEED_VOCAB['size_cap_bytes']


# 扩展名 → kind 映射 (CLI 干跑 + W2 verify 弧机检共用; 运行链主脑显式产 kind
# 时以 spec['kind'] 优先, 本表仅在 kind 缺省时按扩展名兜底猜)。
_EXT_KIND = {
    '.xlsx': 'xlsx', '.csv': 'csv', '.json': 'json', '.jsonl': 'json',
    '.db': 'sqlite', '.sqlite': 'sqlite', '.sqlite3': 'sqlite',
}


def ext_kind_for(path: str) -> str:
    """[COWORK-VERIFY-W2-S1 / 2026-07-09] 按扩展名猜 kind (缺省 'text')。"""
    try:
        return _EXT_KIND.get(os.path.splitext(str(path))[1].lower(), 'text')
    except Exception:
        return 'text'


# ---------------------------------------------------------------------------
# op registry (具名小函数, 准则 10 禁引擎)
#   每个 op: (values: List[Any], expect: Any, spec: dict) -> bool
#   values 由 reader 收集 (已按 ignore_empty 过滤), 保证非空进 op。
# ---------------------------------------------------------------------------

def _eq(a: Any, b: Any) -> bool:
    """相等判定: 直接 == 优先, 退数值容差 (xlsx 数字可能存成 str/float)。"""
    if a == b:
        return True
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return False


def _op_all_equals(values: List[Any], expect: Any, spec: dict) -> bool:
    return all(_eq(v, expect) for v in values)


def _op_equals(values: List[Any], expect: Any, spec: dict) -> bool:
    """单值精确相等 (locator 指向单 cell/单键)。"""
    return len(values) == 1 and _eq(values[0], expect)


def _op_count_gte(values: List[Any], expect: Any, spec: dict) -> bool:
    """非空值个数 ≥ expect (行数/条数下限)。"""
    try:
        return len(values) >= int(expect)
    except (TypeError, ValueError):
        return False


def _op_contains(values: List[Any], expect: Any, spec: dict) -> bool:
    """任一值包含 expect 子串。"""
    needle = str(expect)
    return any(needle in str(v) for v in values)


def _op_nonempty(values: List[Any], expect: Any, spec: dict) -> bool:
    """至少一个非空值 (reader 已过空, 到这里即真非空)。"""
    return any(v not in (None, '') for v in values)


_OPS = {
    'all_equals': _op_all_equals,
    'equals': _op_equals,
    'count_gte': _op_count_gte,
    'contains': _op_contains,
    'nonempty': _op_nonempty,
}


# ---------------------------------------------------------------------------
# 断言解析
# ---------------------------------------------------------------------------

def parse_assertions(raw: Any) -> List[Dict[str, Any]]:
    """JSON str / dict / list → 规整断言 dict list; 坏 → []。

    一个委托挂断言 list = 隐式 AND (全过才 pass, D2)。
    """
    try:
        if isinstance(raw, str):
            s = raw.strip()
            if not s:
                return []
            raw = json.loads(s)
        if isinstance(raw, dict):
            raw = [raw]
        if not isinstance(raw, list):
            return []
        out: List[Dict[str, Any]] = []
        for item in raw:
            if isinstance(item, dict) and item.get('op'):
                out.append(item)
        return out
    except Exception:
        return []


# ---------------------------------------------------------------------------
# reader: _verify_xlsx (S1 唯一; csv/json/sqlite/text 见 W1-S2)
#   签名统一: (spec: dict, artifact_path: str) -> (verdict, evidence)
#   verdict ∈ {'pass', 'fail', 'cannot_tell'}; 任何异常 → cannot_tell。
# ---------------------------------------------------------------------------

def _first_mismatch(cells: List[Tuple[str, Any]], expect: Any, op: str) -> Optional[str]:
    """相等类 op 的首个不符 cell (evidence, 如 'D17=8888')。"""
    if op in ('all_equals', 'equals'):
        for coord, v in cells:
            if not _eq(v, expect):
                return f'{coord}={v!r}'
    return None


def _check_size(path: str) -> bool:
    """size cap 守卫: 超上限返 False (reader → cannot_tell, 防 OOM)。"""
    try:
        return os.path.getsize(path) <= get_size_cap_bytes()
    except OSError:
        return False


def _verify_xlsx(spec: Dict[str, Any], artifact_path: str) -> Tuple[str, Dict[str, Any]]:
    # [COWORK-VERIFY-W1-S3] 前置守卫收敛进 _precheck (policy 红线 + 文件在 + size + op)
    bad = _precheck(spec, artifact_path)
    if bad:
        return bad
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except Exception:
        return ('cannot_tell', {'reason': 'openpyxl_unavailable'})

    locator = spec.get('locator') or {}
    op = spec.get('op')
    expect = spec.get('expect')
    ignore_empty = bool(spec.get('ignore_empty', True))
    sheet = locator.get('sheet', 0)
    column = locator.get('column')

    if column is None:
        return ('cannot_tell', {'reason': 'no_column_locator'})

    wb = None
    try:
        wb = openpyxl.load_workbook(artifact_path, read_only=True, data_only=True)
        names = wb.sheetnames
        if isinstance(sheet, int):
            if sheet < 0 or sheet >= len(names):
                return ('cannot_tell', {'reason': 'sheet_out_of_range'})
            ws = wb[names[sheet]]
        else:
            if sheet not in names:
                return ('cannot_tell', {'reason': 'sheet_not_found'})
            ws = wb[sheet]

        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return ('cannot_tell', {'reason': 'empty_sheet'})

        # 定位列: 支持 header 名 或 int 索引
        col_idx: Optional[int] = None
        if isinstance(column, int):
            col_idx = column
        else:
            for i, h in enumerate(header):
                if h is not None and str(h).strip() == str(column).strip():
                    col_idx = i
                    break
        if col_idx is None:
            return ('cannot_tell', {'reason': f'column_not_found:{column}'})

        col_letter = get_column_letter(col_idx + 1)
        values: List[Any] = []
        cells: List[Tuple[str, Any]] = []
        row_num = 1  # header 占第 1 行
        for row in rows:
            row_num += 1
            if col_idx >= len(row):
                continue
            v = row[col_idx]
            if ignore_empty and (v is None or v == ''):
                continue
            values.append(v)
            cells.append((f'{col_letter}{row_num}', v))

        if not values:
            return ('cannot_tell', {'reason': 'no_values_in_column'})

        ok = _OPS[op](values, expect, spec)
        evidence: Dict[str, Any] = {'rows_checked': len(values), 'op': op, 'expect': expect}
        if not ok:
            evidence['first_mismatch'] = _first_mismatch(cells, expect, op)
        return ('pass' if ok else 'fail', evidence)
    except Exception as e:
        return ('cannot_tell', {'reason': f'read_error:{type(e).__name__}'})
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


# ---------------------------------------------------------------------------
# W1-S2 readers: csv / json / sqlite / text [COWORK-VERIFY-W1-S2 / 2026-07-09]
#   复刻 _verify_xlsx 契约: (spec, path) -> (verdict, evidence);
#   任何异常 → cannot_tell; 空值集 → cannot_tell (防 all_equals 空集空真放松)。
# ---------------------------------------------------------------------------

def _policy_check(path: str) -> Optional[Tuple[str, Dict[str, Any]]]:
    """[COWORK-VERIFY-W1-S3 / 2026-07-09] file_access_policy 红线 (设计 §3.6).

    每 reader 打开文件前必过 check_path(op='read'); fail-toward-DENY:
    policy 拒 / 断供 / 任何异常 → cannot_tell 'policy_denied:...' (只收紧
    不放松, 绝不因 policy 加载失败放行)。过关返 None。
    """
    try:
        from jarvis_file_access_policy import check_path
        allowed, reason = check_path(path, op='read')
        if allowed:
            return None
        return ('cannot_tell', {'reason': f'policy_denied:{reason}'})
    except Exception:
        return ('cannot_tell', {'reason': 'policy_denied:policy_unavailable'})


def _precheck(spec: Dict[str, Any], path: str) -> Optional[Tuple[str, Dict[str, Any]]]:
    """公共前置守卫: policy 红线 + 文件在 + size cap + op 已知; 违规返 cannot_tell 元组, 过关返 None。"""
    denied = _policy_check(path)
    if denied:
        return denied
    if not os.path.isfile(path):
        return ('cannot_tell', {'reason': 'file_not_found'})
    if not _check_size(path):
        return ('cannot_tell', {'reason': 'size_cap_exceeded'})
    if spec.get('op') not in _OPS:
        return ('cannot_tell', {'reason': f"unknown_op:{spec.get('op')}"})
    return None


def _judge(values: List[Any], cells: List[Tuple[str, Any]], spec: Dict[str, Any],
           extra: Optional[Dict[str, Any]] = None) -> Tuple[str, Dict[str, Any]]:
    """收好的非空 values 喂 _OPS 出 verdict + evidence (rows_checked/first_mismatch)。"""
    op = spec['op']
    expect = spec.get('expect')
    ok = _OPS[op](values, expect, spec)
    evidence: Dict[str, Any] = {'rows_checked': len(values), 'op': op, 'expect': expect}
    if extra:
        evidence.update(extra)
    if not ok:
        evidence['first_mismatch'] = _first_mismatch(cells, expect, op)
    return ('pass' if ok else 'fail', evidence)


def _verify_csv(spec: Dict[str, Any], artifact_path: str) -> Tuple[str, Dict[str, Any]]:
    bad = _precheck(spec, artifact_path)
    if bad:
        return bad
    locator = spec.get('locator') or {}
    ignore_empty = bool(spec.get('ignore_empty', True))
    column = locator.get('column')
    if column is None:
        return ('cannot_tell', {'reason': 'no_column_locator'})
    delimiter = str(locator.get('delimiter') or ',')[:1] or ','
    try:
        import csv
        with open(artifact_path, 'r', encoding='utf-8-sig', errors='replace', newline='') as f:
            rows = csv.reader(f, delimiter=delimiter)
            try:
                header = next(rows)
            except StopIteration:
                return ('cannot_tell', {'reason': 'empty_file'})
            col_idx: Optional[int] = None
            if isinstance(column, int):
                col_idx = column
            else:
                for i, h in enumerate(header):
                    if str(h).strip() == str(column).strip():
                        col_idx = i
                        break
            if col_idx is None:
                return ('cannot_tell', {'reason': f'column_not_found:{column}'})
            values: List[Any] = []
            cells: List[Tuple[str, Any]] = []
            row_num = 1  # header 占第 1 行
            for row in rows:
                row_num += 1
                if col_idx >= len(row):
                    continue
                v = row[col_idx]
                if ignore_empty and (v is None or str(v).strip() == ''):
                    continue
                values.append(v)
                cells.append((f'row{row_num}', v))
        if not values:
            return ('cannot_tell', {'reason': 'no_values_in_column'})
        return _judge(values, cells, spec)
    except Exception as e:
        return ('cannot_tell', {'reason': f'read_error:{type(e).__name__}'})


def _json_pluck(node: Any, segments: List[str]) -> List[Any]:
    """按点路径收叶值; list 节点数字段作索引, 非数字段对每项摊平递归。"""
    if not segments:
        return [node]
    seg, rest = segments[0], segments[1:]
    if isinstance(node, dict):
        if seg in node:
            return _json_pluck(node[seg], rest)
        return []
    if isinstance(node, list):
        if seg.lstrip('-').isdigit():
            idx = int(seg)
            if -len(node) <= idx < len(node):
                return _json_pluck(node[idx], segments[1:])
            return []
        out: List[Any] = []
        for item in node:
            out.extend(_json_pluck(item, segments))
        return out
    return []


def _verify_json(spec: Dict[str, Any], artifact_path: str) -> Tuple[str, Dict[str, Any]]:
    bad = _precheck(spec, artifact_path)
    if bad:
        return bad
    locator = spec.get('locator') or {}
    ignore_empty = bool(spec.get('ignore_empty', True))
    key = str(locator.get('key') or '').strip()
    segments = [s for s in key.split('.') if s] if key else []
    try:
        with open(artifact_path, 'r', encoding='utf-8-sig', errors='replace') as f:
            if artifact_path.lower().endswith('.jsonl'):
                docs: List[Any] = []
                for ln, line in enumerate(f, 1):
                    s = line.strip()
                    if not s:
                        continue
                    try:
                        docs.append(json.loads(s))
                    except Exception:
                        return ('cannot_tell', {'reason': f'jsonl_parse_error:line{ln}'})
            else:
                try:
                    docs = [_cfg_overlay.json_load(f)]
                except Exception:
                    return ('cannot_tell', {'reason': 'json_parse_error'})
        values: List[Any] = []
        cells: List[Tuple[str, Any]] = []
        for i, doc in enumerate(docs):
            if not segments and isinstance(doc, list):
                plucked = list(doc)  # 根是 list 且无 key → 摊平 items (count_gte 数条数)
            else:
                plucked = _json_pluck(doc, segments)
            for v in plucked:
                if ignore_empty and (v is None or v == ''):
                    continue
                values.append(v)
                cells.append((f'doc{i}' + (f'.{key}' if key else ''), v))
        if not values:
            reason = f'key_not_found:{key}' if key else 'no_values'
            return ('cannot_tell', {'reason': reason})
        return _judge(values, cells, spec)
    except Exception as e:
        return ('cannot_tell', {'reason': f'read_error:{type(e).__name__}'})


_MEMORY_POOL_DIR = os.path.join(_ROOT, 'memory_pool')
_SQLITE_ROW_CAP = 100000  # system-internal constant (准则6 递归边界)
_SQL_IDENT_RE = re.compile(r'^\w+$')  # 拒引号/分号/空格 — 标识符拼 SQL 防注入


def _sqlite_denied(path: str) -> bool:
    """S2 硬拒 memory_pool 整目录 (真库绝不读; S3 接 file_access_policy 全量策略)。

    判不了 (路径解析异常) 按拒 — fail-toward-DENY。
    """
    try:
        ap = os.path.normcase(os.path.abspath(path))
        mp = os.path.normcase(os.path.abspath(_MEMORY_POOL_DIR)) + os.sep
        return ap.startswith(mp)
    except Exception:
        return True


def _verify_sqlite(spec: Dict[str, Any], artifact_path: str) -> Tuple[str, Dict[str, Any]]:
    if _sqlite_denied(artifact_path):
        return ('cannot_tell', {'reason': 'denied_memory_pool'})
    bad = _precheck(spec, artifact_path)
    if bad:
        return bad
    locator = spec.get('locator') or {}
    op = spec['op']
    expect = spec.get('expect')
    ignore_empty = bool(spec.get('ignore_empty', True))
    table = locator.get('table')
    column = locator.get('column')
    if not table or not _SQL_IDENT_RE.match(str(table)):
        return ('cannot_tell', {'reason': 'bad_table_locator'})
    if column is not None and not _SQL_IDENT_RE.match(str(column)):
        return ('cannot_tell', {'reason': 'bad_column_locator'})
    import sqlite3
    from pathlib import Path
    conn = None
    try:
        uri = Path(artifact_path).resolve().as_uri() + '?mode=ro'  # 只读; 禁 DML (只发 SELECT)
        conn = sqlite3.connect(uri, uri=True, timeout=5)
        cur = conn.cursor()
        if column is None:
            # 行数模式: 只支持 count_gte; 其余 op 必须给 column
            if op != 'count_gte':
                return ('cannot_tell', {'reason': 'no_column_locator'})
            n = int(cur.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0])
            try:
                ok = n >= int(expect)
            except (TypeError, ValueError):
                ok = False  # expect 非数 → fail (与 _op_count_gte 极性一致)
            return ('pass' if ok else 'fail',
                    {'rows_counted': n, 'op': op, 'expect': expect})
        rows = cur.execute(
            f'SELECT "{column}" FROM "{table}" LIMIT {_SQLITE_ROW_CAP + 1}').fetchall()
        if len(rows) > _SQLITE_ROW_CAP:
            return ('cannot_tell', {'reason': 'row_cap_exceeded'})
        values: List[Any] = []
        cells: List[Tuple[str, Any]] = []
        for i, r in enumerate(rows, 1):
            v = r[0]
            if ignore_empty and (v is None or v == ''):
                continue
            values.append(v)
            cells.append((f'{table}.{column}#row{i}', v))
        if not values:
            return ('cannot_tell', {'reason': 'no_values_in_column'})
        return _judge(values, cells, spec)
    except Exception as e:
        return ('cannot_tell', {'reason': f'read_error:{type(e).__name__}'})
    finally:
        if conn is not None:
            try:
                conn.close()
            except Exception:
                pass


def _verify_text(spec: Dict[str, Any], artifact_path: str) -> Tuple[str, Dict[str, Any]]:
    bad = _precheck(spec, artifact_path)
    if bad:
        return bad
    locator = spec.get('locator') or {}
    op = spec['op']
    ignore_empty = bool(spec.get('ignore_empty', True))
    pattern = locator.get('pattern')
    rx = None
    if pattern:
        try:
            rx = re.compile(str(pattern))
        except re.error:
            return ('cannot_tell', {'reason': 'bad_pattern'})
    try:
        with open(artifact_path, 'r', encoding='utf-8', errors='replace') as f:
            all_lines = f.read().splitlines()
    except Exception as e:
        return ('cannot_tell', {'reason': f'read_error:{type(e).__name__}'})
    if ignore_empty:
        lines = [(i, ln) for i, ln in enumerate(all_lines, 1) if ln.strip()]
    else:
        lines = list(enumerate(all_lines, 1))
    if rx is not None:
        matched = [(i, ln) for i, ln in lines if rx.search(ln)]
        extra = {'lines_total': len(all_lines), 'lines_matched': len(matched),
                 'pattern': str(pattern)}
        if not matched:
            # 零命中 = 真观测 (grep 语义): 计数/存在类 op 直接 fail;
            # 相等类无对象可判 → cannot_tell (防 all_equals 空集空真放松)
            if op in ('count_gte', 'nonempty', 'contains'):
                extra.update({'op': op, 'expect': spec.get('expect')})
                return ('fail', extra)
            return ('cannot_tell', {'reason': 'no_matching_lines', **extra})
        return _judge([ln for _, ln in matched],
                      [(f'line{i}', ln) for i, ln in matched], spec, extra=extra)
    if not lines:
        return ('cannot_tell', {'reason': 'empty_file'})
    return _judge([ln for _, ln in lines],
                  [(f'line{i}', ln) for i, ln in lines], spec,
                  extra={'lines_total': len(all_lines)})


_VERIFIERS = {
    'xlsx': _verify_xlsx,
    'csv': _verify_csv,
    'json': _verify_json,
    'sqlite': _verify_sqlite,
    'text': _verify_text,
    # mail 记档不接 (设计 §3.5a: 接线点在 mail organ 非 executor 弧, v1 punt)
}


# ---------------------------------------------------------------------------
# 顶层入口
# ---------------------------------------------------------------------------

def verify_artifact(artifact_path: str, assertions: Any, kind_hint: str = '') -> Dict[str, Any]:
    """按断言机检产物文件, 返 {'verdict', 'checks': [...]}。

    隐式 AND: 任一 fail → fail; 全 pass → pass; 任一 cannot_tell 且无 fail →
    cannot_tell (只收紧不放松)。无断言 → cannot_tell (退兜底, 非 fail)。
    """
    specs = parse_assertions(assertions)
    if not specs:
        return {'verdict': 'cannot_tell', 'checks': [], 'reason': 'no_assertions'}

    checks: List[Dict[str, Any]] = []
    verdicts: List[str] = []
    for spec in specs:
        kind = spec.get('kind') or kind_hint
        verifier = _VERIFIERS.get(kind)
        if verifier is None:
            verdict, evidence = 'cannot_tell', {'reason': f'unknown_kind:{kind}'}
        else:
            try:
                verdict, evidence = verifier(spec, artifact_path)
            except Exception as e:
                verdict, evidence = 'cannot_tell', {'reason': f'verifier_error:{type(e).__name__}'}
        checks.append({
            'kind': kind,
            'op': spec.get('op'),
            'verdict': verdict,
            'evidence': evidence,
        })
        verdicts.append(verdict)

    if 'fail' in verdicts:
        overall = 'fail'
    elif all(v == 'pass' for v in verdicts):
        overall = 'pass'
    else:
        overall = 'cannot_tell'
    return {'verdict': overall, 'checks': checks}
